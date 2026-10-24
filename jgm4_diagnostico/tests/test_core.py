from pathlib import Path
from app.models.database import Base, Parameter, Consultant, Analysis, make_engine
from app.services.database import init_db
from app.services.rule_engine import analyze
from app.services.pricing_service import PricingService
from app.services.import_service import read_text_file
from app.services.manual_chatgpt_service import build_prompt, decode_uploaded_text, parse_response
from app.services.spreadsheet_service import create_collection_workbook
from types import SimpleNamespace

def test_database_and_defaults(tmp_path):
    init_db(tmp_path)
    from sqlalchemy.orm import Session
    with Session(make_engine(tmp_path)) as s:
        assert s.get(Parameter, 'home_rate').value == '120'

def test_rule_analysis_and_schema():
    result = analyze('Custos de fretes em crescimento. Temos transportadoras terceirizadas e ERP sem TMS.')
    assert result.primary_theme == 'Transportes/Fretes'
    assert result.complementary_questions
    assert result.hypotheses[0].source_type == 'INFERENCE'

def test_pricing_reads_rates():
    result = PricingService.calculate(2, 1, 120, 220, expenses=50, discount=10)
    assert result.subtotal == 510
    assert result.final_price == 500

def test_text_import_encodings(tmp_path):
    cases = [('utf8.txt', 'ação logística', 'utf-8'), ('bom.txt', 'ação logística', 'utf-8-sig'), ('win.txt', 'ação logística', 'cp1252'), ('latin.txt', 'ação logística', 'latin-1')]
    for filename, text, encoding in cases:
        path = tmp_path / filename; path.write_bytes(text.encode(encoding)); assert read_text_file(path) == text
    empty = tmp_path / 'empty.txt'; empty.write_bytes(b''); assert read_text_file(empty) == ''

def test_deep_transport_analysis():
    result = analyze('Custos de fretes em crescimento, transportadoras terceirizadas, ERP sem TMS.')
    assert len(result.complementary_questions) >= 10
    assert len(result.hypotheses) >= 5
    assert result.required_data[0].campos_sugeridos

def test_manual_chatgpt_json_and_prompt(tmp_path):
    analysis=SimpleNamespace(company='Empresa Exemplo', theme='Fretes', answers='Custos de frete em alta', consultant_notes='Validar transportadoras')
    result=analyze(analysis.answers, analysis.consultant_notes).model_dump()
    prompt=build_prompt(analysis,result); assert 'Responda somente JSON válido' in prompt
    parsed=parse_response('{"primary_theme":"Transportes/Fretes","recommended_diagnostic":"Diagnóstico de custos","executive_summary":"Resumo","complementary_questions":[{"question":"Qual o gasto?","category":"Custos","order":1}],"required_data":[]}')
    assert parsed.complementary_questions[0].pergunta == 'Qual o gasto?'

def test_custom_collection_workbook(tmp_path):
    analysis=SimpleNamespace(company='Empresa Exemplo')
    path=create_collection_workbook(analysis, {'recommended_diagnostic':'Diagnóstico de fretes','complementary_questions':[],'required_data':[{'campos_sugeridos':['Data','Transportadora']}],'suggested_documents':['Tabela de frete']}, tmp_path)
    assert path.exists()
    from openpyxl import load_workbook
    assert set(load_workbook(path).sheetnames) == {'INSTRUÇÕES','PERGUNTAS','DADOS OPERACIONAIS','DOCUMENTOS SOLICITADOS'}

def test_chatgpt_json_file_variants():
    raw=b'```json\n{"primary_theme":"Fretes","recommended_diagnostic":"Diagnostico","executive_summary":"Resumo","complementary_questions":[],"required_data":[]}\n```'
    assert parse_response(decode_uploaded_text(raw)).primary_theme == 'Fretes'

def test_chatgpt_unknown_source_type_is_normalized():
    raw='{"primary_theme":"Fretes","recommended_diagnostic":"Diagnostico","executive_summary":"Resumo","identified_problems":[{"text":"Problema","source_type":"CHATGPT","confidence":"medium"}],"complementary_questions":[],"required_data":[]}'
    parsed=parse_response(raw)
    assert parsed.identified_problems[0].source_type == 'CLIENT_REPORT'
    assert parsed.identified_problems[0].confidence == 'MEDIUM'

def test_action_plan_pdf_has_logo_and_content(tmp_path):
    from app.services.pdf_service import create_action_plan_pdf
    analysis = SimpleNamespace(company='Empresa Exemplo')
    plan = {'phases': [{
        'order': 1, 'name': 'Levantamento', 'scope': 'Entender a operação',
        'activities': 'Entrevistas; análise preliminar', 'deliverables': 'Plano de coleta',
        'home_hours': 8, 'external_hours': 4,
        'interviews': [{'number': 1, 'participants': 'Direção e logística', 'topics': ['custos', 'prazos']}],
        'required_documents': ['Tabela de fretes', 'Relatório de entregas'],
    }]}
    path = create_action_plan_pdf(analysis, plan, [1], tmp_path)
    assert path.exists() and path.stat().st_size > 1000

def test_action_plan_pdf_normalizes_html_breaks(tmp_path):
    from app.services.pdf_service import _bullet_lines, create_action_plan_pdf
    analysis = SimpleNamespace(company='Empresa Exemplo')
    plan = {'phases': [{'order': 1, 'name': 'Levantamento', 'scope': 'Escopo', 'activities': 'Atividade', 'deliverables': 'Entrega', 'home_hours': 1, 'external_hours': 1, 'interviews': [{'number': 1, 'participants': 'Equipe', 'topics': ['objetivo<br/>impactos']}], 'required_documents': []}]}
    path = create_action_plan_pdf(analysis, plan, [1], tmp_path)
    assert '<br' not in _bullet_lines(['objetivo<br/>impactos'])
    assert path.exists()

def test_contract_uses_only_approved_stages():
    from app.services.contract_service import build_contract_content
    analysis = SimpleNamespace(company='Empresa Exemplo', result={})
    stages = [
        SimpleNamespace(order=1, name='Diagnóstico', description='Escopo 1', home_hours=2, external_hours=1, final_price=500, optional=False),
        SimpleNamespace(order=2, name='Implementação', description='Escopo 2', home_hours=4, external_hours=2, final_price=900, optional=False),
    ]
    content = build_contract_content(analysis, stages, approved_orders={1})
    assert content['approved_orders'] == [1]
    assert [stage['order'] for stage in content['stages']] == [1]
    assert content['total_price'] == 500
