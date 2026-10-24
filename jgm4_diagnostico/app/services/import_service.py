from pathlib import Path
from shutil import copy2
from openpyxl import load_workbook

def read_text_file(file_path: str | Path) -> str:
    """Lê TXT preservando conteúdo e quebras de linha em codificações comuns."""
    path = Path(file_path)
    raw = path.read_bytes()
    if not raw:
        return ''
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeError('Não foi possível decodificar o arquivo TXT')

def import_file(source: str, upload_dir: Path):
    p = Path(source); dest = upload_dir / p.name
    if p.resolve() != dest.resolve(): copy2(p, dest)
    if p.suffix.lower() == '.txt': return dest, {'text': read_text_file(dest)}
    if p.suffix.lower() in ('.xlsx','.xlsm'):
        wb = load_workbook(dest, read_only=True, data_only=True); sheets = {}
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True)); sheets[ws.title] = {'rows': max(0, len(rows)-1), 'columns': [str(x or '') for x in (rows[0] if rows else [])], 'preview': [list(r) for r in rows[:6]]}
        return dest, {'sheets': sheets}
    if p.suffix.lower() == '.csv': return dest, {'text': p.read_text(encoding='utf-8-sig', errors='replace')[:10000]}
    raise ValueError('Formato não suportado')
