from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
def create_collection_workbook(analysis, result, output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True); safe=''.join(c if c.isalnum() else '_' for c in analysis.company); path=output_dir/f'JGM4_{safe}_Coleta_Dados.xlsx'; wb=Workbook(); instructions=wb.active; instructions.title='INSTRUÇÕES'
    instructions.append(['Cliente', analysis.company]); instructions.append(['Projeto', result.get('recommended_diagnostic','Diagnóstico logístico')]); instructions.append(['Objetivo', result.get('executive_summary','')]); instructions.append(['Período solicitado','6 a 12 meses']); instructions.append(['Contato JGM4','JGM4 Consultoria']);
    questions=wb.create_sheet('PERGUNTAS'); questions.append(['Categoria','Pergunta','Resposta','Obrigatória'])
    for q in result.get('complementary_questions',[]): questions.append([q.get('categoria',''),q.get('pergunta',''),'', 'Sim' if q.get('obrigatoria') else 'Não'])
    data=wb.create_sheet('DADOS OPERACIONAIS'); fields=[]
    for item in result.get('required_data',[]): fields.extend(item.get('campos_sugeridos',[]))
    data.append(list(dict.fromkeys(fields or ['Descreva o dado solicitado'])))
    docs=wb.create_sheet('DOCUMENTOS SOLICITADOS'); docs.append(['Documento','Recebido','Observações']); [docs.append([d,'Não','']) for d in result.get('suggested_documents',[])]
    for ws in wb.worksheets:
        for cell in ws[1]: cell.font=Font(bold=True)
        ws.freeze_panes='A2'; ws.column_dimensions['A'].width=28
    wb.save(path); return path
